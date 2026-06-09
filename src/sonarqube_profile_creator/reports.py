from __future__ import annotations

import json
from collections import Counter
from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import reports_dir
from .models import ActionResult, ApplyResult, ItemStatus, PrecheckResult, ProfileSyncPlan, ProfileSyncResult, ValidationIssue


def export_precheck_report(precheck: PrecheckResult, output_dir: Path | None = None) -> tuple[Path, Path]:
    output_dir = output_dir or _timestamped_report_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "precheck_report.json"
    xlsx_path = output_dir / "precheck_report.xlsx"
    json_path.write_text(json.dumps(_serialize_precheck(precheck), indent=2, ensure_ascii=False), encoding="utf-8")
    _write_workbook(precheck, [], xlsx_path)
    return json_path, xlsx_path


def export_apply_report(apply_result: ApplyResult, output_dir: Path | None = None) -> ApplyResult:
    output_dir = output_dir or _timestamped_report_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "report.json"
    xlsx_path = output_dir / "report.xlsx"
    json_path.write_text(json.dumps(_serialize_apply(apply_result), indent=2, ensure_ascii=False), encoding="utf-8")
    _write_workbook(apply_result.precheck, apply_result.actions, xlsx_path)
    apply_result.report_json = json_path
    apply_result.report_xlsx = xlsx_path
    return apply_result


def export_sync_precheck_report(plan: ProfileSyncPlan, output_dir: Path | None = None) -> tuple[Path, Path]:
    output_dir = output_dir or _timestamped_report_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "sync_precheck_report.json"
    xlsx_path = output_dir / "sync_precheck_report.xlsx"
    json_path.write_text(json.dumps(_serialize_sync_plan(plan), indent=2, ensure_ascii=False), encoding="utf-8")
    _write_sync_workbook(plan, [], xlsx_path)
    return json_path, xlsx_path


def export_sync_report(result: ProfileSyncResult, output_dir: Path | None = None) -> ProfileSyncResult:
    output_dir = output_dir or _timestamped_report_dir()
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "sync_report.json"
    xlsx_path = output_dir / "sync_report.xlsx"
    json_path.write_text(json.dumps(_serialize_sync_result(result), indent=2, ensure_ascii=False), encoding="utf-8")
    _write_sync_workbook(result.plan, result.actions, xlsx_path)
    result.report_json = json_path
    result.report_xlsx = xlsx_path
    return result


def _timestamped_report_dir() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return reports_dir() / stamp


def _serialize_precheck(precheck: PrecheckResult) -> dict[str, Any]:
    return {
        "rows": [_dataclass_dict(row) for row in precheck.rows],
        "profile_plans": [_dataclass_dict(plan) for plan in precheck.profile_plans],
        "issues": [_dataclass_dict(issue) for issue in precheck.issues],
        "rule_status": {key: value.value for key, value in precheck.rule_status.items()},
        "default_profiles": precheck.default_profiles,
    }


def _serialize_apply(apply_result: ApplyResult) -> dict[str, Any]:
    return {
        "precheck": _serialize_precheck(apply_result.precheck),
        "actions": [_dataclass_dict(action) for action in apply_result.actions],
        "report_json": str(apply_result.report_json or ""),
        "report_xlsx": str(apply_result.report_xlsx or ""),
    }


def _serialize_sync_plan(plan: ProfileSyncPlan) -> dict[str, Any]:
    return {
        "mode": plan.mode.value,
        "rows": [_dataclass_dict(row) for row in plan.rows],
        "actions": [_dataclass_dict(action) for action in plan.actions],
        "issues": [_dataclass_dict(issue) for issue in plan.issues],
        "profile_keys": {f"{language}:{profile}": key for (language, profile), key in plan.profile_keys.items()},
    }


def _serialize_sync_result(result: ProfileSyncResult) -> dict[str, Any]:
    return {
        "plan": _serialize_sync_plan(result.plan),
        "actions": [_dataclass_dict(action) for action in result.actions],
        "report_json": str(result.report_json or ""),
        "report_xlsx": str(result.report_xlsx or ""),
    }


def _dataclass_dict(value: Any) -> dict[str, Any]:
    if is_dataclass(value):
        payload = asdict(value)
    else:
        payload = dict(value)
    return _json_safe(payload)


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    if isinstance(value, tuple):
        return [_json_safe(item) for item in value]
    if hasattr(value, "value"):
        return value.value
    if isinstance(value, Path):
        return str(value)
    return value


def _write_workbook(precheck: PrecheckResult, actions: list[ActionResult], path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, precheck, actions)
    _write_rows(workbook.create_sheet("Input Rows"), precheck)
    _write_profiles(workbook.create_sheet("Profile Plans"), precheck)
    _write_issues(workbook.create_sheet("Precheck Issues"), precheck.issues)
    _write_actions(workbook.create_sheet("Actions"), actions)
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    workbook.save(path)


def _write_summary(sheet, precheck: PrecheckResult, actions: list[ActionResult]) -> None:
    action_counts = Counter(action.status.value for action in actions)
    issue_counts = Counter(issue.status.value for issue in precheck.issues)
    rows = [
        ("Generated At", datetime.now().isoformat(timespec="seconds")),
        ("Can Apply", "Yes" if precheck.can_apply else "No"),
        ("Input Rows", len(precheck.rows)),
        ("Profile Plans", len(precheck.profile_plans)),
        ("Rules Checked", len(precheck.rule_status)),
        ("Precheck Errors", issue_counts.get(ItemStatus.ERROR.value, 0)),
        ("Precheck Warnings", issue_counts.get(ItemStatus.WARNING.value, 0)),
        ("Action Success", action_counts.get(ItemStatus.OK.value, 0)),
        ("Action Skipped", action_counts.get(ItemStatus.SKIPPED.value, 0)),
        ("Action Errors", action_counts.get(ItemStatus.ERROR.value, 0)),
    ]
    sheet.append(("Metric", "Value"))
    for row in rows:
        sheet.append(row)


def _write_rows(sheet, precheck: PrecheckResult) -> None:
    sheet.append(
        (
            "source_row",
            "language",
            "target_profile",
            "rule_key",
            "parent_profile",
            "strategy",
            "active",
            "sync_action",
            "source_profile",
            "profile_key",
            "rule_name",
            "inheritance",
            "severity",
            "params",
            "prioritizedRule",
            "project_key",
            "set_default",
            "note",
        )
    )
    for row in precheck.rows:
        sheet.append(
            (
                row.source_row,
                row.language,
                row.target_profile,
                row.rule_key,
                row.parent_profile,
                row.strategy,
                row.active,
                row.sync_action,
                row.source_profile,
                row.profile_key,
                row.rule_name,
                row.inheritance,
                row.severity,
                row.params,
                row.prioritizedRule,
                row.project_key,
                row.set_default,
                row.note,
            )
        )


def _write_profiles(sheet, precheck: PrecheckResult) -> None:
    sheet.append(("language", "target_profile", "strategy", "parent_profile", "source_profile", "project_keys", "set_default"))
    for plan in precheck.profile_plans:
        sheet.append(
            (
                plan.language,
                plan.target_profile,
                plan.strategy.value,
                plan.parent_profile,
                plan.source_profile,
                ", ".join(plan.project_keys),
                str(plan.set_default),
            )
        )


def _write_issues(sheet, issues: list[ValidationIssue]) -> None:
    sheet.append(("status", "category", "source_row", "message", "suggestion"))
    for issue in issues:
        sheet.append((issue.status.value, issue.category, issue.source_row, issue.message, issue.suggestion))


def _write_actions(sheet, actions: list[ActionResult]) -> None:
    sheet.append(("status", "action", "language", "profile", "rule_key", "source_row", "message", "api_error", "suggestion"))
    for action in actions:
        sheet.append(
            (
                action.status.value,
                action.action,
                action.language,
                action.profile,
                action.rule_key,
                action.source_row,
                action.message,
                action.api_error,
                action.suggestion,
            )
        )


def _write_sync_workbook(plan: ProfileSyncPlan, actions: list[ActionResult], path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_sync_summary(summary, plan, actions)
    _write_profile_rule_rows(workbook.create_sheet("Profile Rules"), plan)
    _write_issues(workbook.create_sheet("Precheck Issues"), plan.issues)
    _write_actions(workbook.create_sheet("Planned Actions"), plan.actions)
    _write_actions(workbook.create_sheet("Actions"), actions)
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    workbook.save(path)


def _write_sync_summary(sheet, plan: ProfileSyncPlan, actions: list[ActionResult]) -> None:
    action_counts = Counter(action.status.value for action in actions)
    issue_counts = Counter(issue.status.value for issue in plan.issues)
    planned_counts = Counter(action.action for action in plan.actions)
    rows = [
        ("Generated At", datetime.now().isoformat(timespec="seconds")),
        ("Mode", plan.mode.value),
        ("Can Apply", "Yes" if plan.can_apply else "No"),
        ("Input Rows", len(plan.rows)),
        ("Planned Actions", len(plan.actions)),
        ("Plan Activate", planned_counts.get("activate", 0)),
        ("Plan Update", planned_counts.get("update", 0)),
        ("Plan Deactivate", planned_counts.get("deactivate", 0)),
        ("Precheck Errors", issue_counts.get(ItemStatus.ERROR.value, 0)),
        ("Precheck Warnings", issue_counts.get(ItemStatus.WARNING.value, 0)),
        ("Action Success", action_counts.get(ItemStatus.OK.value, 0)),
        ("Action Skipped", action_counts.get(ItemStatus.SKIPPED.value, 0)),
        ("Action Errors", action_counts.get(ItemStatus.ERROR.value, 0)),
    ]
    sheet.append(("Metric", "Value"))
    for row in rows:
        sheet.append(row)


def _write_profile_rule_rows(sheet, plan: ProfileSyncPlan) -> None:
    headers = (
        "source_row",
        "language",
        "target_profile",
        "profile_key",
        "source_profile",
        "rule_key",
        "rule_name",
        "active",
        "severity",
        "params",
        "prioritizedRule",
        "inheritance",
        "sync_action",
        "note",
    )
    sheet.append(headers)
    for row in plan.rows:
        sheet.append(
            (
                row.source_row,
                row.language,
                row.target_profile,
                row.profile_key,
                row.source_profile,
                row.rule_key,
                row.rule_name,
                row.active,
                row.severity,
                row.params,
                row.prioritizedRule,
                row.inheritance,
                row.sync_action,
                row.note,
            )
        )


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = str(cell.value or "")
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 52)
