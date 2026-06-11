from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ALL_FIELDS, FieldMapping, ProfileRuleRow, REQUIRED_FIELDS, RuleRow, SpreadsheetData


PROFILE_RULE_FIELDS = (
    "source_row",
    "language",
    "target_profile",
    "profile_key",
    "source_profile",
    "rule_key",
    "rule_name",
    "active",
    "severity",
    "params",
    "prioritizedRule",
    "inheritance",
    "sync_action",
    "note",
)


HEADER_ALIASES: dict[str, set[str]] = {
    "source_row": {"source_row", "source row", "row", "行号"},
    "language": {"language", "lang", "语言", "语言key", "语言 key", "sonar语言"},
    "target_profile": {
        "target_profile",
        "target profile",
        "quality_profile",
        "quality profile",
        "profile",
        "目标profile",
        "目标 profile",
        "目标质量配置",
        "质量配置",
    },
    "rule_key": {"rule_key", "rule key", "rule", "规则key", "规则 key", "规则", "sonar规则"},
    "parent_profile": {
        "parent_profile",
        "parent profile",
        "parent",
        "父profile",
        "父 profile",
        "继承profile",
        "继承 profile",
    },
    "strategy": {"strategy", "profile_strategy", "策略", "profile策略", "profile 策略"},
    "active": {"active", "enabled", "启用", "是否启用", "激活"},
    "sync_action": {"sync_action", "sync action", "action", "同步动作", "动作"},
    "source_profile": {"source_profile", "source profile", "源profile", "源 profile"},
    "profile_key": {"profile_key", "profile key", "qprofile", "profile id"},
    "rule_name": {"rule_name", "rule name", "name", "规则名称"},
    "inheritance": {"inheritance", "继承状态", "继承"},
    "severity": {"severity", "严重级别", "严重性", "级别"},
    "params": {"params", "parameters", "rule_params", "规则参数", "参数"},
    "prioritizedRule": {
        "prioritizedRule",
        "prioritized_rule",
        "prioritized rule",
        "优先规则",
        "必须修复",
    },
    "project_key": {"project_key", "project key", "project", "项目key", "项目 key", "项目"},
    "set_default": {"set_default", "set default", "default", "设为默认", "默认profile"},
    "note": {"note", "notes", "备注", "说明"},
}


def normalize_header(value: object) -> str:
    text = str(value or "").strip()
    return " ".join(text.replace("-", "_").split()).lower()


def infer_mapping(headers: list[str]) -> FieldMapping:
    normalized = {normalize_header(header): header for header in headers}
    columns: dict[str, str] = {}
    for field in dict.fromkeys((*ALL_FIELDS, *PROFILE_RULE_FIELDS)):
        for alias in HEADER_ALIASES.get(field, set()):
            key = normalize_header(alias)
            if key in normalized:
                columns[field] = normalized[key]
                break
    return FieldMapping(columns=columns)


def read_spreadsheet(path: str | Path, preview_limit: int = 20) -> SpreadsheetData:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        headers, rows = _read_csv(file_path)
    elif suffix in {".xlsx", ".xlsm"}:
        headers, rows = _read_xlsx(file_path)
    else:
        raise ValueError("Supported file types are .csv, .xlsx, and .xlsm.")
    mapping = infer_mapping(headers)
    return SpreadsheetData(
        path=file_path,
        headers=headers,
        preview_rows=rows[:preview_limit],
        rows=rows,
        inferred_mapping=mapping,
    )


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return _read_csv_with_encoding(path, encoding)
        except UnicodeDecodeError:
            continue
    return _read_csv_with_encoding(path, "utf-8-sig")


def _read_csv_with_encoding(path: Path, encoding: str) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding=encoding, newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        dialect = _csv_dialect(sample)
        reader = csv.reader(handle, dialect=dialect)
        raw_headers = next(reader, [])
        headers = _dedupe_headers([str(header or "").strip() for header in raw_headers])
        rows = []
        for raw in reader:
            normalized = _normalize_csv_row(raw, headers)
            if any(normalized.values()):
                rows.append(normalized)
    return headers, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        return [], []
    headers = _dedupe_headers([str(value or "").strip() for value in raw_rows[0]])
    rows: list[dict[str, Any]] = []
    for raw in raw_rows[1:]:
        values = [_clean_value(value) for value in raw]
        if not any(values):
            continue
        rows.append({headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))})
    workbook.close()
    return headers, rows


def _clean_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def rows_from_mapping(rows: list[dict[str, Any]], mapping: FieldMapping) -> list[RuleRow]:
    mapped: list[RuleRow] = []
    for index, raw in enumerate(rows, start=2):
        payload = {}
        for field in ALL_FIELDS:
            source = mapping.source_for(field)
            payload[field] = _clean_value(raw.get(source, "")) if source else ""
        mapped.append(
            RuleRow(
                source_row=index,
                language=payload["language"],
                target_profile=payload["target_profile"],
                rule_key=payload["rule_key"],
                parent_profile=payload["parent_profile"],
                strategy=payload["strategy"],
                active=payload["active"],
                sync_action=payload["sync_action"],
                source_profile=payload["source_profile"],
                profile_key=payload["profile_key"],
                rule_name=payload["rule_name"],
                inheritance=payload["inheritance"],
                severity=payload["severity"],
                params=payload["params"],
                prioritizedRule=payload["prioritizedRule"],
                project_key=payload["project_key"],
                set_default=payload["set_default"],
                note=payload["note"],
                raw=raw,
            )
        )
    return mapped


def profile_rule_rows_from_mapping(rows: list[dict[str, Any]], mapping: FieldMapping) -> list[ProfileRuleRow]:
    mapped: list[ProfileRuleRow] = []
    for index, raw in enumerate(rows, start=2):
        payload = {}
        for field in PROFILE_RULE_FIELDS:
            source = mapping.source_for(field)
            payload[field] = _clean_value(raw.get(source, "")) if source else ""
        mapped.append(
            ProfileRuleRow(
                source_row=_parse_source_row(payload["source_row"], index),
                language=payload["language"],
                target_profile=payload["target_profile"],
                profile_key=payload["profile_key"],
                source_profile=payload["source_profile"],
                rule_key=payload["rule_key"],
                rule_name=payload["rule_name"],
                active=payload["active"] or "true",
                severity=payload["severity"],
                params=payload["params"],
                prioritizedRule=payload["prioritizedRule"],
                inheritance=payload["inheritance"],
                sync_action=payload["sync_action"],
                note=payload["note"],
                raw=raw,
            )
        )
    return mapped


def _csv_dialect(sample: str) -> csv.Dialect:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        return csv.excel


def _dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers, start=1):
        base = header or f"column_{index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        result.append(base if count == 1 else f"{base}_{count}")
    return result


def _normalize_csv_row(row: list[Any], headers: list[str]) -> dict[str, str]:
    return {header: _clean_value(row[index]) if index < len(row) else "" for index, header in enumerate(headers)}


def _parse_source_row(value: object, fallback: int) -> int:
    text = _clean_value(value)
    if not text:
        return fallback
    try:
        return int(text)
    except ValueError:
        try:
            number = float(text)
        except ValueError:
            return fallback
        return int(number) if number.is_integer() else fallback


def write_profile_rules(output_dir: str | Path, rows: list[ProfileRuleRow], name: str = "profile_rules") -> tuple[Path, Path]:
    path = Path(output_dir)
    path.mkdir(parents=True, exist_ok=True)
    csv_path = path / f"{name}.csv"
    xlsx_path = path / f"{name}.xlsx"
    payloads = [_profile_rule_payload(row) for row in rows]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(PROFILE_RULE_FIELDS))
        writer.writeheader()
        writer.writerows(payloads)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Profile Rules"
    sheet.append(list(PROFILE_RULE_FIELDS))
    for payload in payloads:
        sheet.append([payload[field] for field in PROFILE_RULE_FIELDS])
    _style_sheet(sheet)
    workbook.save(xlsx_path)
    return csv_path, xlsx_path


def validate_required_mapping(mapping: FieldMapping, required_fields: tuple[str, ...] = REQUIRED_FIELDS) -> list[str]:
    return [field for field in required_fields if not mapping.source_for(field)]


def write_example_templates(output_dir: str | Path) -> tuple[Path, Path]:
    path = Path(output_dir)
    path.mkdir(parents=True, exist_ok=True)
    rows = [
        {
            "rule_key": "java:S1144",
            "active": "true",
            "severity": "MAJOR",
            "params": "",
            "prioritizedRule": "false",
            "note": "",
        },
        {
            "rule_key": "javascript:S1128",
            "active": "true",
            "severity": "",
            "params": "",
            "prioritizedRule": "false",
            "note": "",
        },
    ]
    csv_path = path / "sonarqube_rules_template.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = path / "sonarqube_rules_template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rules"
    sheet.append(list(rows[0].keys()))
    for row in rows:
        sheet.append(list(row.values()))
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 14), 36)
    workbook.save(xlsx_path)
    return csv_path, xlsx_path


def _profile_rule_payload(row: ProfileRuleRow) -> dict[str, Any]:
    return {
        "source_row": row.source_row,
        "language": row.language,
        "target_profile": row.target_profile,
        "profile_key": row.profile_key,
        "source_profile": row.source_profile,
        "rule_key": row.rule_key,
        "rule_name": row.rule_name,
        "active": row.active,
        "severity": row.severity,
        "params": row.params,
        "prioritizedRule": row.prioritizedRule,
        "inheritance": row.inheritance,
        "sync_action": row.sync_action,
        "note": row.note,
    }


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = str(cell.value or "")
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 52)
